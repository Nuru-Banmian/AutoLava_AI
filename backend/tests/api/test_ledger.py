from dataclasses import dataclass
import asyncio
from io import BytesIO
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

import app.models.agent  # noqa: F401 - register all tables before the session schema fixture

from app.models.identity import Store, StoreMember, User
from app.models.ledger import IncomeCategory, StoreDailyRecord
from app.services.scheduler import apply_refreshed_weather
from app.services.weather import WeatherResult
from app.services.weather import RECORD_WEATHER_OPTIONS


@dataclass
class AssignedStore:
    store: Store
    cash: IncomeCategory
    excluded: IncomeCategory
    store_id: int
    cash_id: int
    excluded_id: int
    timezone: str

    @property
    def id(self) -> int:
        return self.store_id


@pytest.fixture
async def assigned_store(
    auth_client: AsyncClient, db_session: AsyncSession, store_factory
) -> AssignedStore:
    user = await db_session.scalar(select(User).where(User.username == "authenticated"))
    assert user is not None
    store = await store_factory(name="Assigned", timezone="Europe/Berlin")
    store.income_items_enabled = True
    cash = IncomeCategory(
        store_id=store.id, name="Cash", include_in_total=True, is_active=True, sort_order=0
    )
    excluded = IncomeCategory(
        store_id=store.id,
        name="Excluded",
        include_in_total=False,
        is_active=True,
        sort_order=1,
    )
    db_session.add_all([StoreMember(store_id=store.id, user_id=user.id), cash, excluded])
    await db_session.commit()
    return AssignedStore(
        store,
        cash,
        excluded,
        store_id=store.id,
        cash_id=cash.id,
        excluded_id=excluded.id,
        timezone=store.timezone,
    )


@pytest.fixture
def ledger_payload(assigned_store: AssignedStore) -> dict:
    return {
        "is_open": "营业",
        "daily_revenue": None,
        "wash_count": 12,
        "weather": "晴",
        "weather_edited": True,
        "activity": None,
        "items": [
            {"category_id": assigned_store.cash_id, "amount": 200},
            {"category_id": assigned_store.excluded_id, "amount": 80},
        ],
    }


def today_for(assigned_store: AssignedStore) -> date:
    return datetime.now(ZoneInfo(assigned_store.timezone)).date()


async def put_ledger(client: AsyncClient, path: str, *, json: dict):
    current = await client.get(path)
    expected = current.json() if current.status_code == 200 else None
    store_id, record_date = path.rsplit("/", 2)[-2:]
    form = (await client.get(f"/api/ledger/{store_id}/{record_date}/form-config")).json()
    return await client.put(path, json=json | {
        "expected_identity": expected["identity"] if expected else None,
        "expected_revision": expected["revision"] if expected else None,
        "expected_config_revision": form["config_revision"],
    })


async def delete_ledger(client: AsyncClient, path: str):
    current = (await client.get(path)).json()
    return await client.request("DELETE", path, json={
        "expected_identity": current["identity"],
        "expected_revision": current["revision"],
    })


async def test_unreported_can_save_only_status_and_preserve_unknown_values(
    auth_client: AsyncClient, assigned_store: AssignedStore,
) -> None:
    path = f"/api/ledger/{assigned_store.id}/2026-07-01"
    response = await put_ledger(auth_client, path, json={"is_open": "未统计"})
    assert response.status_code == 201, response.text
    assert response.json()["daily_revenue"] is None
    record = (await auth_client.get(path)).json()
    assert record["is_open"] == "未统计"
    assert record["daily_revenue"] is None
    assert record["wash_count"] is None
    assert record["items"] == []


async def test_unreported_clear_requires_fresh_confirmation_and_keeps_context(
    auth_client: AsyncClient, assigned_store: AssignedStore, ledger_payload: dict,
) -> None:
    path = f"/api/ledger/{assigned_store.id}/2026-07-02"
    original = await put_ledger(auth_client, path, json=ledger_payload | {"activity": "集中清点前三天收入"})
    assert original.status_code == 201
    before = (await auth_client.get(path)).json()
    denied = await put_ledger(auth_client, path, json={"is_open": "未统计"})
    assert denied.status_code == 409
    assert denied.json()["detail"]["code"] == "clear_values_confirmation_required"
    assert denied.json()["detail"]["current"]["daily_revenue"] == 200
    assert (await auth_client.get(path)).json()["revision"] == before["revision"]
    cleared = await put_ledger(auth_client, path, json={"is_open": "未统计", "confirm_clear_values": True})
    assert cleared.status_code == 200
    after = (await auth_client.get(path)).json()
    assert (after["id"], after["identity"], after["income_mode"]) == (before["id"], before["identity"], before["income_mode"])
    assert after["revision"] == before["revision"] + 1
    assert (after["weather"], after["activity"]) == ("晴", "集中清点前三天收入")
    assert (after["daily_revenue"], after["wash_count"], after["items"]) == (None, None, [])
    stale = await auth_client.put(path, json={
        "is_open": "未统计", "confirm_clear_values": True,
        "expected_identity": before["identity"], "expected_revision": before["revision"],
        "expected_config_revision": before["config_revision"],
    })
    assert stale.status_code == 409
    assert stale.json()["detail"]["code"] == "ledger_revision_conflict"
    assert (await auth_client.get(path)).json()["revision"] == after["revision"]
    back = await put_ledger(auth_client, path, json={"is_open": "营业"})
    assert back.status_code == 422
    assert (await auth_client.get(path)).json()["is_open"] == "未统计"
    restored = await put_ledger(auth_client, path, json=ledger_payload)
    assert restored.status_code == 200


@pytest.mark.parametrize("values", [
    {"daily_revenue": 0}, {"wash_count": 0}, {"items": [{"category_id": 1, "amount": 0}]},
])
async def test_unreported_rejects_even_zero_values(
    auth_client: AsyncClient, assigned_store: AssignedStore, values: dict,
) -> None:
    path = f"/api/ledger/{assigned_store.id}/2026-07-03"
    response = await put_ledger(auth_client, path, json={"is_open": "未统计"} | values)
    assert response.status_code == 422
    assert (await auth_client.get(path)).status_code == 404


async def test_unreported_return_requires_explicit_enabled_quantity(
    auth_client: AsyncClient, assigned_store: AssignedStore, ledger_payload: dict,
) -> None:
    path = f"/api/ledger/{assigned_store.id}/2026-07-08"
    assert (await put_ledger(auth_client, path, json={"is_open": "未统计"})).status_code == 201
    for count in (None,):
        denied = await put_ledger(auth_client, path, json=ledger_payload | {"wash_count": count})
        assert denied.status_code == 422
    omitted = await put_ledger(auth_client, path, json={k: v for k, v in ledger_payload.items() if k != "wash_count"})
    assert omitted.status_code == 422
    assert (await auth_client.get(path)).json()["is_open"] == "未统计"
    saved = await put_ledger(auth_client, path, json=ledger_payload | {"wash_count": 0})
    assert saved.status_code == 200
    assert (await auth_client.get(path)).json()["wash_count"] == 0


async def test_unreported_preserves_omitted_historical_weather_and_can_clear_context(
    auth_client: AsyncClient, assigned_store: AssignedStore, ledger_payload: dict,
    db_session: AsyncSession,
) -> None:
    path = f"/api/ledger/{assigned_store.id}/2026-07-09"
    assert (await put_ledger(auth_client, path, json=ledger_payload | {"activity": "原事件"})).status_code == 201
    record = await db_session.scalar(select(StoreDailyRecord).where(StoreDailyRecord.store_id == assigned_store.id))
    record.weather = "旧版任意天气"
    await db_session.commit()
    clear_values = await put_ledger(auth_client, path, json={"is_open": "未统计", "confirm_clear_values": True})
    assert clear_values.status_code == 200
    saved = (await auth_client.get(path)).json()
    assert (saved["weather"], saved["activity"], saved["daily_revenue"]) == ("旧版任意天气", "原事件", None)
    clear_context = await put_ledger(auth_client, path, json={"is_open": "未统计", "weather": None, "activity": None})
    assert clear_context.status_code == 200
    saved = (await auth_client.get(path)).json()
    assert (saved["weather"], saved["activity"]) == (None, None)


async def test_unreported_analysis_retains_state_but_excludes_unknown_values(
    auth_client: AsyncClient, assigned_store: AssignedStore, ledger_payload: dict,
) -> None:
    base = f"/api/ledger/{assigned_store.id}"
    for day, payload in [
        (1, ledger_payload), (2, ledger_payload | {"is_open": "提前休息"}),
        (3, ledger_payload | {"is_open": "休息"}),
        (4, {"is_open": "未统计", "activity": "尚未清点"}),
        (5, ledger_payload | {"items": [{"category_id": assigned_store.cash_id, "amount": 0}, {"category_id": assigned_store.excluded_id, "amount": 0}]}),
    ]:
        assert (await put_ledger(auth_client, f"{base}/2026-07-0{day}", json=payload)).status_code == 201
    assert (await put_ledger(auth_client, f"{base}/2026-06-01", json={"is_open": "未统计"})).status_code == 201
    response = await auth_client.get(f"/api/charts/{assigned_store.id}", params={
        "start": "2026-07-01", "end": "2026-07-06", "compare_start": "2026-06-01", "compare_end": "2026-06-06",
    })
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["kpis"]["total_revenue"] == 400
    assert data["kpis"]["open_days"] == 3
    assert data["kpis"]["average_revenue"] == 133
    assert data["period_coverage"] == {
        "start": "2026-07-01", "end": "2026-07-06", "record_days": 5, "interval_days": 6,
        "statistical_days": 4, "unreported_days": 1, "operating_days": 3, "rest_days": 1, "missing_record_days": 1,
    }
    assert data["daily"][3] == {"date": "2026-07-04", "revenue": None, "is_open": "未统计"}
    assert data["ledger_comparison"]["status"] == "no_previous_records"
    assert data["ledger_comparison"]["change_percent"] is None
    all_unknown = (await auth_client.get(f"/api/charts/{assigned_store.id}?start=2026-06-01&end=2026-06-06")).json()
    assert all_unknown["monthly"] == [{"month": "2026-06", "revenue": None, "daily_ledger_revenue": None, "confirmed_settlement_income": 0, "monthly_total_income": None}]
    page = (await auth_client.get(f"/api/database/{assigned_store.id}/records?status=未统计")).json()
    assert page["total"] == 2
    missing = (await auth_client.get(f"/api/database/{assigned_store.id}/records?missing_wash_count=true")).json()
    assert all(row["is_open"] in ("营业", "提前休息") for row in missing["items"])
    exported = await auth_client.get(f"/api/database/{assigned_store.id}/export.xlsx")
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.content), read_only=True)
    unknown_row = next(row for row in workbook["经营记录"].values if "尚未清点" in row)
    assert unknown_row[2:5] == ("未统计", None, None)
    assert unknown_row[6] == "尚未清点"
    assert not any(row[0] and str(row[0]).startswith("2026-07-04") for row in list(workbook["收入明细"].values)[1:])


async def test_unreported_cached_briefings_and_recent_keep_unknown_state(
    auth_client: AsyncClient, assigned_store: AssignedStore,
) -> None:
    today = today_for(assigned_store)
    for day in (today - timedelta(days=1), today):
        assert (await put_ledger(auth_client, f"/api/ledger/{assigned_store.id}/{day}", json={"is_open": "未统计"})).status_code == 201
    refreshed = await auth_client.post(f"/api/dashboard/{assigned_store.id}/refresh")
    assert refreshed.status_code == 200, refreshed.text
    cached = await auth_client.get(f"/api/dashboard/{assigned_store.id}")
    assert cached.status_code == 200
    cards = {card["card_type"]: card for card in cached.json()}
    for name in ("today", "yesterday"):
        assert cards[name]["state"] == "unreported"
        assert cards[name]["revenue"] is None
    recent = (await auth_client.get(f"/api/ledger/{assigned_store.id}/recent")).json()
    assert len(recent) == 2
    assert all(row["is_open"] == "未统计" and row["daily_revenue"] is None for row in recent)


async def test_clear_confirmation_for_real_zero_checks_config_before_confirmation(
    auth_client: AsyncClient, assigned_store: AssignedStore, ledger_payload: dict,
) -> None:
    path = f"/api/ledger/{assigned_store.id}/2026-07-07"
    zero_items = [{"category_id": item["category_id"], "amount": 0} for item in ledger_payload["items"]]
    assert (await put_ledger(auth_client, path, json=ledger_payload | {"items": zero_items})).status_code == 201
    current = (await auth_client.get(path)).json()
    stale_config = await auth_client.put(path, json={
        "is_open": "未统计", "confirm_clear_values": True,
        "expected_identity": current["identity"], "expected_revision": current["revision"],
        "expected_config_revision": 99,
    })
    assert stale_config.status_code == 409
    assert stale_config.json()["detail"]["code"] == "income_config_revision_conflict"
    absent = await put_ledger(auth_client, path, json={"is_open": "未统计"})
    assert absent.status_code == 409
    assert absent.json()["detail"]["code"] == "clear_values_confirmation_required"
    assert (await auth_client.get(path)).json()["daily_revenue"] == 0


async def test_put_releases_dependency_transaction_before_weather(
    auth_client: AsyncClient,
    assigned_store: AssignedStore,
    ledger_payload: dict,
    db_session: AsyncSession,
) -> None:
    observed_transactions: list[bool] = []

    class TransactionObservingWeather:
        async def get_daily(self, store: Store, target: date):
            observed_transactions.append(db_session.in_transaction())
            return None

    auth_client._transport.app.state.weather_service = TransactionObservingWeather()

    response = await put_ledger(auth_client,
        f"/api/ledger/{assigned_store.id}/{today_for(assigned_store).isoformat()}",
        json=ledger_payload,
    )

    assert response.status_code == 201
    assert observed_transactions == []


async def test_historical_ledger_save_and_read_do_not_wait_for_weather(
    auth_client: AsyncClient,
    assigned_store: AssignedStore,
    ledger_payload: dict,
) -> None:
    entered = asyncio.Event()
    release = asyncio.Event()

    class BlockedWeather:
        async def get_daily(self, store: Store, target: date):
            entered.set()
            await release.wait()
            return WeatherResult("晴", 0, 20.0, 10.0, 0.0)

    auth_client._transport.app.state.weather_service = BlockedWeather()
    record_date = today_for(assigned_store) - timedelta(days=30)
    path = f"/api/ledger/{assigned_store.id}/{record_date.isoformat()}"
    try:
        response = await asyncio.wait_for(
            put_ledger(auth_client, path, json=ledger_payload), timeout=1
        )
        assert response.status_code == 201
        assert (await auth_client.get(path)).json()["daily_revenue"] == 200
    finally:
        release.set()


@pytest.mark.parametrize("amount", [1.5, "1.00"])
async def test_amount_input_requires_json_integer(
    auth_client: AsyncClient,
    assigned_store: AssignedStore,
    ledger_payload: dict,
    amount,
) -> None:
    ledger_payload["items"][0]["amount"] = amount
    response = await put_ledger(auth_client,
        f"/api/ledger/{assigned_store.id}/{today_for(assigned_store).isoformat()}",
        json=ledger_payload,
    )
    assert response.status_code == 422


@pytest.mark.parametrize("daily_revenue", [1.5, "1.00"])
async def test_direct_total_requires_json_integer(
    auth_client: AsyncClient,
    assigned_store: AssignedStore,
    db_session: AsyncSession,
    daily_revenue,
) -> None:
    assigned_store.store.income_items_enabled = False
    await db_session.commit()
    response = await put_ledger(auth_client,
        f"/api/ledger/{assigned_store.id}/{today_for(assigned_store).isoformat()}",
        json={"is_open": "营业", "daily_revenue": daily_revenue, "items": []},
    )
    assert response.status_code == 422


async def test_second_put_overwrites_without_compatibility_parameters(
    auth_client: AsyncClient, assigned_store: AssignedStore, ledger_payload: dict
) -> None:
    path = f"/api/ledger/{assigned_store.id}/{today_for(assigned_store).isoformat()}"
    first = await put_ledger(auth_client, path, json=ledger_payload)
    second_payload = ledger_payload | {
        "items": [
            {"category_id": assigned_store.cash_id, "amount": 321},
            {"category_id": assigned_store.excluded_id, "amount": 90},
        ]
    }
    second = await put_ledger(auth_client, path, json=second_payload)
    assert first.status_code == 201
    assert second.status_code == 200
    assert second.json()["daily_revenue"] == 321


async def test_record_card_returns_persisted_bookkeeping_events_only_on_database_view(
    auth_client: AsyncClient,
    assigned_store: AssignedStore,
    ledger_payload: dict,
    db_session: AsyncSession,
) -> None:
    path = f"/api/ledger/{assigned_store.id}/{today_for(assigned_store).isoformat()}"

    created = await put_ledger(auth_client, path, json=ledger_payload)
    updated = await put_ledger(auth_client,
        path,
        json=ledger_payload
        | {
            "items": [
                {"category_id": assigned_store.cash_id, "amount": 321},
                {"category_id": assigned_store.excluded_id, "amount": 90},
            ]
        },
    )
    saved_record = await db_session.scalar(
        select(StoreDailyRecord).where(StoreDailyRecord.store_id == assigned_store.id)
    )
    assert saved_record is not None
    apply_refreshed_weather(
        saved_record,
        WeatherResult("多云", 3, 24.0, 16.0, 0.0),
    )
    await db_session.commit()
    card_records = await auth_client.get(
        f"/api/database/{assigned_store.id}/records"
    )
    direct_record = await auth_client.get(path)

    assert created.status_code == 201
    assert updated.status_code == 200
    assert card_records.status_code == 200
    events = card_records.json()["items"][0]["bookkeeping_events"]
    assert [event["action"] for event in events] == ["created", "updated"]
    assert [event["actor_name"] for event in events] == [
        "authenticated",
        "authenticated",
    ]
    assert len({event["id"] for event in events}) == 2
    assert [event["timestamp_status"] for event in events] == ["utc", "utc"]
    assert all(
        datetime.fromisoformat(event["occurred_at"].replace("Z", "+00:00")).utcoffset()
        == timedelta(0)
        for event in events
    )
    assert "bookkeeping_events" not in direct_record.json()


async def test_record_snapshot_is_retained_after_current_category_edits(
    auth_client: AsyncClient,
    assigned_store: AssignedStore,
    ledger_payload: dict,
    db_session: AsyncSession,
) -> None:
    path = f"/api/ledger/{assigned_store.id}/{today_for(assigned_store).isoformat()}"
    assert (await put_ledger(auth_client, path, json=ledger_payload)).status_code == 201
    assigned_store.cash.name = "Renamed"
    assigned_store.cash.include_in_total = False
    assigned_store.cash.sort_order = 8
    assigned_store.excluded.name = "Excluded renamed"
    assigned_store.excluded.include_in_total = True
    assigned_store.excluded.sort_order = 9
    await db_session.commit()

    updated = await put_ledger(auth_client,
        path,
        json=ledger_payload
        | {
            "items": [
                {"category_id": assigned_store.cash_id, "amount": 125},
                {"category_id": assigned_store.excluded_id, "amount": 75},
            ]
        },
    )
    fetched = await auth_client.get(path)
    assert updated.status_code == 200
    assert updated.json()["daily_revenue"] == 125
    assert [
        (item["category_name"], item["include_in_total"], item["sort_order"])
        for item in fetched.json()["items"]
    ] == [("Cash", True, 0), ("Excluded", False, 1)]


async def test_existing_record_accepts_a_newly_enabled_income_category(
    auth_client: AsyncClient,
    assigned_store: AssignedStore,
    ledger_payload: dict,
    db_session: AsyncSession,
) -> None:
    path = f"/api/ledger/{assigned_store.id}/{today_for(assigned_store).isoformat()}"
    assert (await put_ledger(auth_client, path, json=ledger_payload)).status_code == 201
    added = IncomeCategory(
        store_id=assigned_store.id,
        name="Added later",
        include_in_total=True,
        is_active=True,
        sort_order=2,
    )
    db_session.add(added)
    await db_session.flush()
    added_id = added.id
    await db_session.commit()

    form_config = await auth_client.get(f"{path}/form-config")
    assert [item["category_id"] for item in form_config.json()["items"]] == [
        assigned_store.cash_id,
        assigned_store.excluded_id,
        added_id,
    ]

    updated = await put_ledger(auth_client,
        path,
        json=ledger_payload
        | {
            "items": [
                *ledger_payload["items"],
                {"category_id": added_id, "amount": 30},
            ]
        },
    )

    assert updated.status_code == 200
    assert updated.json()["daily_revenue"] == 230
    fetched = await auth_client.get(path)
    assert [(item["category_id"], item["amount"]) for item in fetched.json()["items"]] == [
        (assigned_store.cash_id, 200),
        (assigned_store.excluded_id, 80),
        (added_id, 30),
    ]


async def test_put_and_get_return_integer_money(
    auth_client: AsyncClient, assigned_store: AssignedStore, ledger_payload: dict
) -> None:
    path = f"/api/ledger/{assigned_store.id}/{today_for(assigned_store).isoformat()}"
    created = await put_ledger(auth_client, path, json=ledger_payload)
    fetched = await auth_client.get(path)
    assert created.status_code == 201
    assert created.json()["daily_revenue"] == 200
    assert fetched.json()["daily_revenue"] == 200
    assert [item["amount"] for item in fetched.json()["items"]] == [200, 80]


async def test_early_close_preserves_values_while_rest_normalizes_them(
    auth_client: AsyncClient, assigned_store: AssignedStore, ledger_payload: dict
) -> None:
    path = f"/api/ledger/{assigned_store.id}/{today_for(assigned_store).isoformat()}"

    early_close = await put_ledger(auth_client,
        path,
        json=ledger_payload | {"is_open": "提前休息"},
    )
    early_close_record = await auth_client.get(path)

    assert early_close.status_code == 201
    assert early_close_record.json()["is_open"] == "提前休息"
    assert early_close_record.json()["daily_revenue"] == 200
    assert early_close_record.json()["wash_count"] == 12
    assert [item["amount"] for item in early_close_record.json()["items"]] == [200, 80]

    rest = await put_ledger(auth_client, path, json=ledger_payload | {"is_open": "休息"})
    rest_record = await auth_client.get(path)

    assert rest.status_code == 200
    assert rest_record.json()["is_open"] == "休息"
    assert rest_record.json()["daily_revenue"] == 0
    assert rest_record.json()["wash_count"] == 0
    assert [item["amount"] for item in rest_record.json()["items"]] == [0, 0]


async def test_recent_uses_store_local_window(
    auth_client: AsyncClient, assigned_store: AssignedStore, ledger_payload: dict
) -> None:
    today = today_for(assigned_store)
    for target in (today - timedelta(days=7), today - timedelta(days=2), today):
        response = await put_ledger(auth_client,
            f"/api/ledger/{assigned_store.id}/{target.isoformat()}", json=ledger_payload
        )
        assert response.status_code == 201
    recent = await auth_client.get(
        f"/api/ledger/{assigned_store.id}/recent", params={"days": 7}
    )
    assert [item["date"] for item in recent.json()] == [
        today.isoformat(),
        (today - timedelta(days=2)).isoformat(),
    ]


async def test_future_and_invalid_status_are_422(
    auth_client: AsyncClient, assigned_store: AssignedStore, ledger_payload: dict
) -> None:
    future = await put_ledger(auth_client,
        f"/api/ledger/{assigned_store.id}/2999-01-01", json=ledger_payload
    )
    invalid = await put_ledger(auth_client,
        f"/api/ledger/{assigned_store.id}/{today_for(assigned_store).isoformat()}",
        json=ledger_payload | {"is_open": "unknown"},
    )
    legacy = await put_ledger(auth_client,
        f"/api/ledger/{assigned_store.id}/{today_for(assigned_store).isoformat()}",
        json=ledger_payload | {"is_open": "天气停业"},
    )
    assert future.status_code == invalid.status_code == legacy.status_code == 422


async def test_delete_returns_204(
    auth_client: AsyncClient,
    assigned_store: AssignedStore,
    ledger_payload: dict,
    db_session: AsyncSession,
) -> None:
    user = await db_session.scalar(select(User).where(User.username == "authenticated"))
    assert user is not None
    user.role = "admin"
    await db_session.commit()
    path = f"/api/ledger/{assigned_store.id}/{today_for(assigned_store).isoformat()}"
    assert (await put_ledger(auth_client, path, json=ledger_payload)).status_code == 201
    deleted = await delete_ledger(auth_client, path)
    assert deleted.status_code == 204
    assert (await auth_client.get(path)).status_code == 404


async def test_record_weather_contract_rejects_new_unknown_values_without_writing(
    auth_client: AsyncClient, assigned_store: AssignedStore, ledger_payload: dict,
) -> None:
    target = today_for(assigned_store).isoformat()
    path = f"/api/ledger/{assigned_store.id}/{target}"
    options = await auth_client.get("/api/ledger/weather-options")
    assert options.status_code == 200
    assert options.json() == list(RECORD_WEATHER_OPTIONS)
    assert len(options.json()) == 28
    for label in options.json():
        assert (await put_ledger(auth_client, path, json=ledger_payload | {"weather": label})).status_code in (200, 201)
    assert (await put_ledger(auth_client, path, json=ledger_payload | {"weather": None})).status_code == 200
    invalid = await put_ledger(auth_client, path, json=ledger_payload | {
        "weather": "供应商未知天气",
        "items": [{"category_id": assigned_store.cash_id, "amount": 999},
                  {"category_id": assigned_store.excluded_id, "amount": 80}],
    })
    assert invalid.status_code == 422
    saved = (await auth_client.get(path)).json()
    assert saved["weather"] is None
    assert saved["daily_revenue"] == 200


async def test_legacy_weather_reads_exports_and_analysis_until_corrected(
    auth_client: AsyncClient, assigned_store: AssignedStore, ledger_payload: dict,
    db_session: AsyncSession,
) -> None:
    target = today_for(assigned_store).isoformat()
    path = f"/api/ledger/{assigned_store.id}/{target}"
    assert (await put_ledger(auth_client, path, json=ledger_payload)).status_code == 201
    record = await db_session.scalar(select(StoreDailyRecord).where(StoreDailyRecord.store_id == assigned_store.id))
    assert record is not None
    record.weather = "旧版任意天气"
    record.weather_edited = False
    await db_session.commit()

    read = await auth_client.get(path)
    assert read.status_code == 200
    assert read.json()["weather"] == "旧版任意天气"
    assert read.json()["weather_legacy"] is True
    listing = await auth_client.get(f"/api/database/{assigned_store.id}/records")
    assert listing.json()["items"][0]["weather_legacy"] is True
    chart = await auth_client.get(f"/api/charts/{assigned_store.id}", params={"start": target, "end": target})
    assert chart.status_code == 200
    assert chart.json()["weather"] == [
        {"weather": "历史未规范天气", "average_revenue": 200, "operating_day_count": 1}
    ]
    exported = await auth_client.get(f"/api/database/{assigned_store.id}/export.xlsx")
    assert exported.status_code == 200
    rows = load_workbook(BytesIO(exported.content), read_only=True)["经营记录"].values
    assert "历史旧值：旧版任意天气" in next(row for row in rows if target in str(row[0]))
    assert (await put_ledger(auth_client, path, json=ledger_payload | {"weather": "旧版任意天气"})).status_code == 422
    assert (await put_ledger(auth_client, path, json={key: value for key, value in ledger_payload.items() if key != "weather"})).status_code == 422
    assert (await put_ledger(auth_client, path, json=ledger_payload | {"weather": None, "weather_edited": True})).status_code == 200
    corrected = (await auth_client.get(path)).json()
    assert corrected["weather"] is None
    assert corrected["weather_legacy"] is False
    assert corrected["weather_edited"] is True
    apply_refreshed_weather(record, WeatherResult("晴", 0, 20.0, 10.0, 0.0))
    await db_session.commit()
    after_refresh = (await auth_client.get(path)).json()
    assert after_refresh["weather"] is None
    assert after_refresh["weather_auto"] == "晴"


@pytest.mark.parametrize("cached", [False, True])
async def test_explicit_clear_with_auto_weather_stays_cleared(
    auth_client: AsyncClient, assigned_store: AssignedStore, ledger_payload: dict,
    db_session: AsyncSession, cached: bool,
) -> None:
    target = today_for(assigned_store).isoformat()
    path = f"/api/ledger/{assigned_store.id}/{target}"
    assert (await put_ledger(auth_client, path, json=ledger_payload)).status_code == 201
    record = await db_session.scalar(select(StoreDailyRecord).where(StoreDailyRecord.store_id == assigned_store.id))
    assert record is not None
    record.weather = None
    record.weather_auto = "晴" if cached else None
    record.weather_edited = False
    await db_session.commit()

    class FreshWeather:
        async def get_daily(self, store, requested_date):
            return WeatherResult("晴", 0, 20.0, 10.0, 0.0)

    auth_client._transport.app.state.weather_service = FreshWeather()

    clear = await put_ledger(auth_client, path, json=ledger_payload | {"weather": None, "weather_edited": False})
    assert clear.status_code == 200
    saved = (await auth_client.get(path)).json()
    assert saved["weather"] is None
    assert saved["weather_edited"] is True


async def test_second_save_does_not_wait_for_fresh_weather(
    auth_client: AsyncClient, assigned_store: AssignedStore, ledger_payload: dict,
) -> None:
    path = f"/api/ledger/{assigned_store.id}/{today_for(assigned_store).isoformat()}"
    body = {key: value for key, value in ledger_payload.items() if key != "weather"}
    body["weather_edited"] = False
    assert (await put_ledger(auth_client, path, json=body)).status_code == 201

    class FreshWeather:
        async def get_daily(self, store, requested_date):
            return WeatherResult("晴", 0, 20.0, 10.0, 0.0)

    auth_client._transport.app.state.weather_service = FreshWeather()
    assert (await put_ledger(auth_client, path, json=body)).status_code == 200
    saved = (await auth_client.get(path)).json()
    assert saved["weather"] is None
    assert saved["weather_edited"] is False
